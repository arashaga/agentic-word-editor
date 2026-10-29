"""Validate cited research and create editable DOCX reports using python-docx."""

from datetime import datetime, timezone
from io import BytesIO
import re
from urllib.parse import urlsplit

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from pydantic import BaseModel, ConfigDict

from app.documents import INVALID_XML, MAX_AI_CHARACTERS, WordDocument
from app.errors import DemoError


class ResearchSource(BaseModel):
    """A provider-returned citation title and URL, validated before report creation."""

    model_config = ConfigDict(extra="forbid")
    title: str
    url: str


class ReportSection(BaseModel):
    """Plain-text section content with one-based references into the source list."""

    model_config = ConfigDict(extra="forbid")
    heading: str
    paragraphs: list[str]
    source_ids: list[int]


class ResearchReport(BaseModel):
    """Structured writer output used to construct a new, styled Word document."""

    model_config = ConfigDict(extra="forbid")
    title: str
    subtitle: str
    sections: list[ReportSection]


def valid_source(source: ResearchSource) -> bool:
    """Check citation length, XML-safe single-line text, and a credential-free HTTPS URL."""
    try:
        url = urlsplit(source.url)
        return (
            url.scheme == "https" and bool(url.hostname) and not url.username and not url.password
            and url.port in (None, 443) and len(source.url) <= 2048
            and not INVALID_XML.search(source.title + source.url)
            and not any(c in source.title + source.url for c in "\r\n\t")
            and 0 < len(source.title) <= 500
        )
    except ValueError:
        return False


def create_report(report: ResearchReport, sources: list[ResearchSource]) -> WordDocument:
    """Create a US Letter DOCX with editable sections and numbered source references.

    Validate text, citation indices, and size limits before returning a WordDocument.
    Raise DemoError for invalid report data or an unusable template. This checks
    citation structure, not factual accuracy, and never saves the report to OneDrive.
    """
    if not 1 <= len(sources) <= 30 or not all(valid_source(source) for source in sources):
        raise DemoError("Research must include 1-30 valid HTTPS source citations.", 422)
    if not 1 <= len(report.sections) <= 12:
        raise DemoError("A report must contain 1-12 sections.", 422)
    texts = [report.title, report.subtitle]
    for section in report.sections:
        if not 1 <= len(section.paragraphs) <= 8 or not section.source_ids:
            raise DemoError("Each report section needs paragraphs and source citations.", 422)
        if any(source_id < 1 or source_id > len(sources) for source_id in section.source_ids):
            raise DemoError("The writer cited a source that was not returned by web search.", 422)
        texts.extend([section.heading, *section.paragraphs])
    if any(not text.strip() or INVALID_XML.search(text) or any(c in text for c in "\r\n\t") for text in texts):
        raise DemoError("The report contains empty or unsupported text. Generate a new report.", 422)
    if len(report.title) > 160 or sum(map(len, texts)) + sum(len(s.title + s.url) for s in sources) > 24000:
        raise DemoError("The research report exceeds the demo's document size limit.", 413)
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        if not isinstance(style, ParagraphStyle):
            raise DemoError("The Word report template is missing a paragraph style.", 500)
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string("000000")
        if name == "Normal":
            style.font.size = Pt(11)
    doc.add_heading(report.title, 0)
    doc.add_paragraph(report.subtitle)
    doc.add_paragraph("Researched " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    for item in report.sections:
        doc.add_heading(item.heading, 1)
        for paragraph in item.paragraphs:
            doc.add_paragraph(paragraph)
        doc.add_paragraph("Sources: " + ", ".join(f"[{number}]" for number in dict.fromkeys(item.source_ids)))
    doc.add_heading("Sources", 1)
    for number, source in enumerate(sources, 1):
        doc.add_paragraph(f"[{number}] {source.title}")
        doc.add_paragraph(source.url)
    doc.add_paragraph("AI-assisted research. Check the cited releases and final Word layout before relying on this report.")
    buffer = BytesIO()
    doc.save(buffer)
    stem = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "-", report.title).strip(" .")[:100] or "Research report"
    result = WordDocument(buffer.getvalue(), stem + ".docx")
    if sum(len(block.text) for block in result.blocks) > MAX_AI_CHARACTERS:
        raise DemoError("The generated document exceeds the editing limit.", 413)
    return result
